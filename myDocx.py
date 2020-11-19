import docx
import re
import os

from configuration import Configuration


class MyDocx:
    config = Configuration().doc_config_dict

    def __init__(self, document_path):
        self.document = docx.Document(document_path)

    def get_text(self):
        text = self.document.text
        return text

    def get_table_list(self):
        table_list = self.document.tables
        return table_list

    @staticmethod
    def _get_table_all_rows(table):
        rows = table.rows
        return rows

    @staticmethod
    def _get_table_all_columns(table):
        columns = table.columns
        return columns

    @staticmethod
    def get_table_cell(table, row=None, column=None):
        if row is not None and column is not None:

            ret_cell = table.cell(row_idx=row, col_idx=column).text
            # all_rows = MyDocx._get_table_all_rows(table)
            # ret_cell = all_rows[row].cells[column].text
        elif row is not None and column is None:
            ret_cell = list()
            all_rows = MyDocx._get_table_all_rows(table)
            for cell in all_rows[row].cells:
                ret_cell.append(cell.text)
        elif row is None and column is not None:
            ret_cell = list()
            all_columns = MyDocx._get_table_all_columns(table)
            for cell in all_columns[column].cells:
                ret_cell.append(cell.text)
        else:
            ret_cell = None
        return ret_cell

    @staticmethod
    def check_table_sql(table):
        check_flag = MyDocx.get_table_cell(table=table,
                                           row=int(MyDocx.config['tableFlagPosition']['row']),
                                           column=int(MyDocx.config['tableFlagPosition']['col']))
        # print(table._column_count)
        # cell_idx = 1 + (-2 * table._column_count)
        # print(cell_idx)
        # for cell in MyDocx.get_table_cell(table=table, row=0):
        #     print(cell)

        if check_flag.startswith(MyDocx.config['tableFlagValue']):
            # print('1')
            return True
        print('config.tableFlagValue(数据字典标识)：', check_flag, ' 检测错误，与config内配置值：', MyDocx.config['tableFlagValue'], ' 不匹配')
        return False

    @staticmethod
    def check_repeated_table(table_list):
        table_name_set = set()
        repeated_table_name = set()
        table_name_list = list()
        for table in table_list:
            table_name = MyDocx.get_table_cell(table=table, row=int(MyDocx.config['tableNamePosition']['row']), column=int(MyDocx.config['tableNamePosition']['col']))
            table_name_set.add(table_name)
            table_name_list.append(table_name)
            if len(table_name_list) != len(table_name_set):
                repeated_table_name.add(table_name)
                table_name_list.pop(-1)
        return list(repeated_table_name)

    @staticmethod
    def get_sql_table_list(table_list):
        sql_table_list = list()
        count = 0
        for table in table_list:
            if len(MyDocx._get_table_all_rows(table)) < int(MyDocx.config['minTableRowsValue']):
                continue
            count += 1
            try:
                if MyDocx.check_table_sql(table=table):
                    sql_table_list.append(table)
            except:
                print('Error in table:' + str(MyDocx.get_table_cell(table=table, row=-3, column=1)) + u'表识别错误')
        return sql_table_list

    @staticmethod
    def get_table_info(table):
        if table is None:
            return None
        flag = True
        row = 1
        table_info = dict()
        # print(table)
        table_name = MyDocx.get_table_cell(table=table,
                                           row=int(MyDocx.config['tableNamePosition']['row']),
                                           column=int(MyDocx.config['tableNamePosition']['col']))
        table_info['table_name'] = str(table_name).replace(' ', '')
        
        table_info['params'] = dict()
        column_names = MyDocx.get_table_cell(table, row=0)
        name_num = 0
        type_num = 0
        null_num = 0
        only_num = 0
        index_num = 0
        comment_num = 0
        current_column = 0
        default_num = 0
        for column_name in column_names:
            if column_name.startswith(u'字段名'):
                name_num = current_column
            if column_name.startswith(u'字段类型'):
                type_num = current_column
            if column_name.startswith(u'空') or column_name.startswith(u'为空'):
                null_num = current_column
            if column_name.startswith(u'唯一'):
                only_num = current_column
            if column_name.startswith(u'索引'):
                index_num = current_column
            if column_name.startswith(u'COMMENT') or column_name.startswith(u'注释'):
                comment_num = current_column
            if column_name.startswith(u'默认值'):
                default_num = current_column
            current_column += 1
        try:
            while flag:
                param_map = dict()
                param_map['type'] = ''
                name = MyDocx.get_table_cell(table=table, row=row, column=name_num)

                type_word_list = MyDocx.get_table_cell(table=table, row=row, column=type_num).replace('\n', '').replace('\r', '').lower().split(' ')

                unsigned_flag = False
                for type_word in type_word_list:
                    if type_word == 'unsigned':
                        unsigned_flag = True
                        continue
                    else:
                        param_map['type'] += type_word
                if unsigned_flag:
                    param_map['type'] += ' unsigned'
                param_map['type_ori'] = MyDocx.get_table_cell(table=table, row=row, column=type_num)
                is_null = MyDocx.get_table_cell(table=table, row=row, column=null_num).replace(' ', '').lower()
                if is_null == '是' or is_null == 'y' or is_null == 'yes':
                    param_map['is_null'] = 'y'
                else:
                    param_map['is_null'] = 'n'
                param_map['is_null_ori'] = MyDocx.get_table_cell(table=table, row=row, column=null_num)
                is_only = MyDocx.get_table_cell(table=table, row=row, column=only_num).replace(' ', '')
                if is_only == '是' or is_only == 'y' or is_only == 'yes':
                    param_map['is_only'] = 'y'
                else:
                    param_map['is_only'] = 'n'
                param_map['is_only_ori'] = MyDocx.get_table_cell(table=table, row=row, column=only_num)
                is_index = MyDocx.get_table_cell(table=table, row=row, column=index_num).replace(' ', '').lower()
                if is_index == '是' or is_index == 'y' or is_index == 'yes' or is_index == '主键':
                    param_map['is_index'] = 'y'
                elif is_index == 'pri' or is_index == 'uni' or is_index == 'mul':
                    param_map['is_index'] = is_index.upper()
                else:
                    param_map['is_index'] = 'n'
                param_map['is_index_ori'] = MyDocx.get_table_cell(table=table, row=row, column=index_num)

                param_map['comment_ori'] = MyDocx.get_table_cell(table=table, row=row, column=comment_num)
                if param_map['comment_ori'] is None:
                    param_map['comment_ori'] = ''
                else:
                    # param_map['comment'] = param_map['comment_ori'].replace(' ', '').replace('\n', '').replace('\r', '')
                    param_map['comment'] = param_map['comment_ori'].replace('\n', '\\n').replace('\r', '\\r').replace('\r\n', '\\r\\n')
                    # param_map['comment'] = param_map['comment_ori']
                    # print(param_map['comment'])

                if default_num != 0:
                    param_map['default_ori'] = MyDocx.get_table_cell(table=table, row=row, column=default_num)
                    param_map['default'] = param_map['default_ori']
                else:
                    param_map['default_ori'] = ''
                    param_map['default'] = 'NULL'

                table_info['params'][str(name)] = param_map
                row += 1
                # print(table_name)
                # print(name_num)
                # print(MyDocx.get_table_cell(table=table, row=row, column=name_num))
                # print(MyDocx.get_table_cell(table=table, row=row, column=0))
                if MyDocx.get_table_cell(table=table, row=row,
                                         column=int(MyDocx.config['fieldInfoEndColumn'])).startswith(MyDocx.config['fieldInfoEndValue']):
                    flag = False
            if MyDocx.config['tableEngineFlag']:
                table_info['database_engine_ori'] = MyDocx.get_table_cell(table=table,
                                                                          row=int(MyDocx.config['tableEnginePosition']['row']),
                                                                          column=int(MyDocx.config['tableEnginePosition']['col']))
                table_info['database_engine'] = table_info['database_engine_ori'].lower()

            if MyDocx.config['tableDefaultCharsetFlag']:
                table_info['default_charset_ori'] = MyDocx.get_table_cell(table=table,
                                                                          row=int(MyDocx.config['tableDefaultCharsetPosition']['row']),
                                                                          column=int(MyDocx.config['tableDefaultCharsetPosition']['col']))
                table_info['default_charset'] = table_info['default_charset_ori'].lower()

            if MyDocx.config['unionIndexFlag']:
                table_info['union_index_ori'] = MyDocx.get_table_cell(table=table,
                                                                      row=int(MyDocx.config['unionIndexPosition']['row']),
                                                                      column=int(MyDocx.config['unionIndexPosition']['col']))
                table_info['union_index'] = dict()
                if len(table_info['union_index_ori']) > 3:
                    key_name_list = re.findall('([A-z_0-9]+)\(', table_info['union_index_ori'])
                    key_pos = 0
                    for key_name in key_name_list:
                        table_info['union_index'][key_name] = list()
                        params_str_lists = re.findall('\(([A-z_,0-9 ]+)\)', table_info['union_index_ori'])
                        # print(params_str_lists)
                        table_info['union_index'][key_name] = str(params_str_lists[key_pos]).replace(' ', '').split(',')
                        # print(table_info['union_index'][key_name])
                        key_pos += 1
        except Exception as err:
            print(str(table_name) + u'数据读取错误:' + str(err))
            table_info = None
        # print(table_info)
        # if table_name == 'badge':
        #     print(table_info)
        return table_info

    def get_info_list(self):
        sql_table_list = self.get_sql_table_list(self.get_table_list())
        repeated_table = self.check_repeated_table(sql_table_list)
        print(u'存在重复的表名：', repeated_table)
        info_list = list()
        for sql_table in sql_table_list:
            # print(sql_table)
            table_info = self.get_table_info(sql_table)
            info_list.append(table_info)
        return info_list


if __name__ == '__main__':
    my = MyDocx('')
    # my = MyDocx('')
    tl = my.get_info_list()
    for t in tl:
        print(t)




